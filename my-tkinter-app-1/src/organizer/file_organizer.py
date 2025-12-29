import sys
import os
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..')))
def move_files(regrouped, all_files):
    """
    Déplace les fichiers selon la structure {theme: {sous_theme: [files]}}.
    all_files : liste de dicts avec 'name' et 'path'.
    """
    import shutil
    import os
    # Création d'un mapping nom -> chemin
    file_map = {f['name']: f['path'] for f in all_files}
    for theme, sous_dict in regrouped.items():
        for sous_theme, files_list in sous_dict.items():
            for file_name in files_list:
                src = file_map.get(file_name)
                if not src or not os.path.exists(src):
                    continue
                dest_dir = os.path.join(os.path.dirname(src), theme)
                if sous_theme:
                    dest_dir = os.path.join(dest_dir, sous_theme)
                os.makedirs(dest_dir, exist_ok=True)
                dest = os.path.join(dest_dir, file_name)
                try:
                    shutil.move(src, dest)
                except Exception as e:
                    print(f"Erreur lors du déplacement de {src} vers {dest} : {e}")


import os
import json
from ai.gemini_validator import GeminiValidator
from datetime import datetime

try:
    import docx
except ImportError:
    docx = None
try:
    import openpyxl
except ImportError:
    openpyxl = None

SCHEMA_PATH = 'src/data/schema.json'

def extract_text_excerpt(filepath, max_chars=300):
    ext = os.path.splitext(filepath)[1].lower()
    try:
        if ext in ['.txt', '.csv']:
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read(max_chars)
        elif ext == '.json':
            import json
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                try:
                    data = json.load(f)
                    text = json.dumps(data, ensure_ascii=False)
                    return text[:max_chars]
                except Exception:
                    f.seek(0)
                    return f.read(max_chars)
        elif ext == '.docx' and docx:
            doc = docx.Document(filepath)
            text = '\n'.join([p.text for p in doc.paragraphs])
            return text[:max_chars]
        elif ext in ['.xlsx', '.xls'] and openpyxl:
            wb = openpyxl.load_workbook(filepath, read_only=True)
            ws = wb.active
            content = []
            for row in ws.iter_rows(values_only=True):
                content.append(' | '.join([str(cell) for cell in row if cell is not None]))
                if len(' '.join(content)) > max_chars:
                    break
            return ' '.join(content)[:max_chars]
    except Exception:
        return ''
    return ''

def get_all_files(directory):
    files = []
    for root, dirs, filenames in os.walk(directory):
        # Exclure les dossiers .git
        dirs[:] = [d for d in dirs if d != '.git']
        for file in filenames:
            path = os.path.join(root, file)
            ext = os.path.splitext(file)[1].lower()
            if ext in ['.txt', '.csv', '.json', '.docx', '.xlsx', '.xls']:
                excerpt = extract_text_excerpt(path)
            else:
                excerpt = ''
            try:
                ctime = os.path.getctime(path)
                date_str = datetime.fromtimestamp(ctime).strftime('%Y-%m-%d %H:%M:%S')
            except Exception:
                date_str = ''
            files.append({
                'name': file,
                'path': path,
                'date': date_str,
                'excerpt': excerpt
            })
    return files



def get_default_user_dirs():
    """Retourne les chemins Documents, Téléchargements et Bureau de l'utilisateur Windows."""
    from pathlib import Path
    import os
    home = Path.home()
    docs = os.path.join(home, 'Documents')
    downloads = os.path.join(home, 'Downloads')
    desktop = os.path.join(home, 'Desktop')
    return [d for d in [docs, downloads, desktop] if os.path.isdir(d)]

def organize_files(directory=None):
    dirs_to_scan = [directory] if directory else get_default_user_dirs()
    all_files = []
    for d in dirs_to_scan:
        all_files.extend(get_all_files(d))

    # Charger ou générer le schéma d'organisation par sujets
    if not os.path.exists(SCHEMA_PATH):
        gemini = GeminiValidator()
        organization_schema = gemini.suggest_schema(all_files)
        with open(SCHEMA_PATH, 'w', encoding='utf-8') as schema_file:
            json.dump(organization_schema, schema_file, indent=2, ensure_ascii=False)
    else:
        with open(SCHEMA_PATH, 'r', encoding='utf-8') as schema_file:
            organization_schema = json.load(schema_file)

    print("Test de move_files avec l'organisation détectée...")
    regrouped = {}
    for file in all_files:
        filename = file['name']
        subject = organization_schema.get(filename)
        if subject and isinstance(subject, dict):
            theme = subject.get('theme', 'Divers')
            sous_theme = subject.get('sous_theme', '')
            if theme not in regrouped:
                regrouped[theme] = {}
            if sous_theme not in regrouped[theme]:
                regrouped[theme][sous_theme] = []
            regrouped[theme][sous_theme].append(filename)
    print("Organisation regroupée :", json.dumps(regrouped, indent=2, ensure_ascii=False))
    move_files(regrouped, all_files)
    print("Déplacement terminé.")


def main():
    organize_files()

if __name__ == "__main__":
    main()